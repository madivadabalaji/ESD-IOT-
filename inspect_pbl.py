import docx

doc = docx.Document(r'c:\Users\subha\Downloads\Telegram Desktop\PBL Final Documentation templet (3).docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")

for i in range(30):
    p = doc.paragraphs[i]
    print(f"P{i:02d} [{p.style.name}]: {repr(p.text)}")
